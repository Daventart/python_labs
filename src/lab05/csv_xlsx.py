"""Конвертация CSV -> XLSX."""

import csv
from pathlib import Path

from openpyxl import Workbook
from openpyxl.utils import get_column_letter

from src.lib.io_helpers import ensure_parent_dir


def csv_to_xlsx(csv_path: str, xlsx_path: str) -> None:
    
    src = Path(csv_path)
    dst = Path(xlsx_path)

    if not src.exists():
        raise FileNotFoundError(f"CSV-файл не найден: {src}")
    if src.suffix.lower() != ".csv":
        raise ValueError(f"Ожидался файл .csv, получен: {src.name}")
    if dst.suffix.lower() != ".xlsx":
        raise ValueError(f"Ожидался файл .xlsx, получен: {dst.name}")

    with src.open("r", newline="", encoding="utf-8") as f:
        rows = list(csv.reader(f))

    if not rows:
        raise ValueError("CSV пустой")

    wb = Workbook()
    ws = wb.active
    ws.title = "Sheet1"

    for row in rows:
        ws.append(row)

    widths: dict[int, int] = {}
    for row in rows:
        for idx, value in enumerate(row, start=1):
            widths[idx] = max(widths.get(idx, 0), len(value))

    for idx, width in widths.items():
        ws.column_dimensions[get_column_letter(idx)].width = max(width, 8)

    ensure_parent_dir(dst)
    wb.save(str(dst))


if __name__ == "__main__":
    csv_to_xlsx("data/samples/people.csv", "data/out/people.xlsx")
    csv_to_xlsx("data/samples/cities.csv", "data/out/cities.xlsx")
    print("CSV -> XLSX: готово")